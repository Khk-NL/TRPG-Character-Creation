#!/usr/bin/env python3
"""Fill the bundled CoC 7e workbook without rebuilding its OOXML package."""

from __future__ import annotations

import argparse
import ast
import difflib
import hashlib
import json
import math
import os
import random
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")


MAIN_SHEET = "人物卡"
NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL_DOC = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_REL_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
ET.register_namespace("", NS_MAIN)
ET.register_namespace("r", NS_REL_DOC)

CHAR_KEYS = ("STR", "CON", "SIZ", "DEX", "APP", "INT", "POW", "EDU", "LUCK")
CHAR_CELLS = {
    "STR": "U3", "CON": "U5", "SIZ": "U7",
    "DEX": "AA3", "APP": "AA5", "INT": "AA7",
    "POW": "AG3", "EDU": "AG5", "LUCK": "AG7",
}
BACKGROUND_CELLS = {
    "description": "AA61",
    "ideology": "AA63",
    "significant_person": "AA65",
    "meaningful_place": "AA67",
    "treasured_possession": "AA69",
    "trait": "AA71",
    "story": "W79",
}
OPEN_CHOICE_CELLS = ["BA18", "BA19", "BA20", "BA21", "BA22", "BJ18", "BJ19"]
SPECIALTY_ALIASES = {
    "艺术及手艺": "技艺", "艺术手艺": "技艺", "艺术和手艺": "技艺",
    "艺术与手艺": "技艺", "技艺": "技艺",
    "外语": "其他语言", "语言": "其他语言", "其他语言": "其他语言",
    "操纵": "驾驶", "驾驶": "驾驶", "格斗": "格斗", "射击": "射击",
    "科学": "科学", "生存": "生存", "学识": "学识", "学问": "学识",
}
WEAPON_ALIASES = {
    "小刀": "小型刀具", "折刀": "小型刀具", "随身刀": "小型刀具",
    "野外小刀": "小型刀具", "匕首": "小型刀具",
    "砍刀": "大型刀具", "柴刀": "大型刀具", "开山刀": "大型刀具",
}
CHINESE_NUMBERS = {"一": 1, "二": 2, "两": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}


class PlanError(ValueError):
    pass


def norm(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = text.replace("（", "(").replace("）", ")").replace("：", ":")
    return re.sub(r"[\s·•,，.。:：()（）/\\_-]+", "", text)


def as_int(value: Any, label: str) -> int:
    if isinstance(value, bool):
        raise PlanError(f"{label} 必须是整数")
    try:
        number = int(value)
    except (TypeError, ValueError) as exc:
        raise PlanError(f"{label} 必须是整数") from exc
    return number


def output_filename(name: str, group_name: str) -> str:
    def safe_part(value: str, label: str) -> str:
        cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "-", str(value or "")).strip(" .-")
        if not cleaned:
            raise PlanError(f"生成文件名需要{label}；请填写 basic.group_name 或从公告确认团名")
        return cleaned
    return f"{safe_part(name, '角色名')}-{safe_part(group_name, '团名')}.xlsx"


def split_specialty(name: str) -> tuple[str, str | None]:
    text = str(name).strip().replace("（", "(").replace("）", ")").replace("：", ":")
    match = re.match(r"^\s*([^(:]+?)\s*(?:\(\s*([^)]+?)\s*\)|:\s*(.+))\s*$", text)
    if not match:
        return text, None
    return match.group(1).strip(), (match.group(2) or match.group(3)).strip()


def workbook_catalog(template: Path) -> dict[str, Any]:
    wb = load_workbook(template, read_only=True, data_only=False, keep_links=True)
    try:
        if MAIN_SHEET not in wb.sheetnames or "职业列表" not in wb.sheetnames or "武器列表" not in wb.sheetnames:
            raise PlanError("模板缺少人物卡、职业列表或武器列表页")

        occupations = []
        ws = wb["职业列表"]
        for row, cells in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
            serial, name = cells[0], cells[1]
            if not name:
                continue
            occupations.append({
                "row": row,
                "serial": serial,
                "name": str(name).strip(),
                "credit_min": cells[3],
                "credit_max": cells[4],
                "points_label": cells[5],
                "points_formula": cells[6],
                "skills": cells[7],
                "description": cells[13],
            })

        skills = []
        card = wb[MAIN_SHEET]
        for row, cells in enumerate(card.iter_rows(min_row=16, max_row=49, values_only=True), start=16):
            for side, primary_col, specialty_col, occ_col, int_col in (
                ("left", "F", "H", "N", "P"),
                ("right", "AB", "AD", "AJ", "AL"),
            ):
                primary = cells[column_number(primary_col + "1") - 1]
                specialty = cells[column_number(specialty_col + "1") - 1]
                if not primary:
                    continue
                display = str(primary).strip()
                if specialty:
                    display += f"({str(specialty).strip()})"
                skills.append({
                    "row": row, "side": side, "name": display,
                    "primary": str(primary).strip(),
                    "specialty": str(specialty).strip() if specialty else "",
                    "specialty_cell": f"{specialty_col}{row}",
                    "occupation_cell": f"{occ_col}{row}",
                    "interest_cell": f"{int_col}{row}",
                })

        open_choice_names = {
            str(wb["附表"][f"X{row}"].value).strip()
            for row in range(25, 93) if wb["附表"][f"X{row}"].value is not None
        }

        weapons = []
        weapon_sheet = wb["武器列表"]
        for row, cells in enumerate(weapon_sheet.iter_rows(min_row=2, values_only=True), start=2):
            name = cells[1]
            if not name:
                continue
            weapons.append({
                "row": row,
                "name": str(name).strip(),
                "skill": cells[2],
                "damage": cells[3],
                "range": cells[4],
            })
        return {
            "template": str(template),
            "sheets": wb.sheetnames,
            "occupations": occupations,
            "skills": skills,
            "open_choice_names": sorted(open_choice_names),
            "weapons": weapons,
        }
    finally:
        wb.close()


def weapon_candidates(query: str, weapons: list[dict[str, Any]], limit: int = 5) -> list[str]:
    normalized = norm(query)
    expanded = norm(WEAPON_ALIASES.get(normalized, query))
    def score(weapon: dict[str, Any]) -> tuple[float, float]:
        name = norm(weapon["name"])
        similarity = max(difflib.SequenceMatcher(None, normalized, name).ratio(),
                         difflib.SequenceMatcher(None, expanded, name).ratio())
        if expanded and expanded in name:
            similarity += 1
        if normalized and normalized in name:
            similarity += 0.5
        return similarity, -len(name)
    ranked = sorted(weapons, key=score, reverse=True)
    return [str(item["name"]) for item in ranked[:limit]]


def open_occupation_slots(description: str) -> list[dict[str, Any]]:
    """Summarize explicit free-choice slots; do not mistake them for interest skills."""
    source = str(description or "")
    patterns = (
        r"(?:任意|任选|自选)?\s*([一二两三四五六七八九十\d]+)\s*项\s*其他[^，。；;]*?(?:特长|技能|专业|领域)",
        r"(?:任意|任选|自选)\s*([一二两三四五六七八九十\d]+)\s*项[^，。；;]*?(?:特长|技能|专业|领域)",
        r"(?:任意|任选|自选)\s*其他\s*([一二两三四五六七八九十\d]+)\s*项[^，。；;]*?(?:特长|技能|专业)",
        r"([一二两三四五六七八九十\d]+)\s*个\s*学习的专业",
    )
    found = []
    spans: list[tuple[int, int]] = []
    for pattern in patterns:
        for match in re.finditer(pattern, source):
            if any(a <= match.start() < b for a, b in spans):
                continue
            token = match.group(1)
            count = int(token) if token.isdigit() else CHINESE_NUMBERS.get(token)
            if count is not None:
                found.append({"count": count, "description": match.group(0)})
                spans.append(match.span())
    return sorted(found, key=lambda item: source.find(item["description"]))


def resolve_named(query: str, records: list[dict[str, Any]], label: str) -> dict[str, Any]:
    q = norm(query)
    if label == "武器" and q in WEAPON_ALIASES:
        target = norm(WEAPON_ALIASES[q])
        aliases = [r for r in records if target in norm(r.get("name"))]
        if len(aliases) == 1:
            return aliases[0]
    exact = [r for r in records if norm(r.get("name")) == q]
    if len(exact) == 1:
        return exact[0]
    partial = [r for r in records if q and (q in norm(r.get("name")) or norm(r.get("name")) in q)]
    if len(partial) == 1:
        return partial[0]
    choices = exact or partial
    if choices:
        names = "、".join(str(r.get("name")) for r in choices[:8])
        raise PlanError(f"{label}“{query}”匹配到多个条目：{names}")
    if label == "武器":
        suggestions = "、".join(weapon_candidates(query, records))
        raise PlanError(f"模板中找不到武器“{query}”；是否指：{suggestions}")
    raise PlanError(f"模板中找不到{label}“{query}”")


def resolve_skill(query: str, skills: list[dict[str, Any]], claimed: set[str]) -> tuple[dict[str, Any], str | None]:
    q = norm(query)
    exact = [s for s in skills if norm(s["name"]) == q]
    for skill in exact:
        if skill["occupation_cell"] not in claimed:
            return skill, None

    parent, specialty = split_specialty(query)
    if specialty:
        wanted_parent = SPECIALTY_ALIASES.get(norm(parent), parent)
        candidates = []
        for skill in skills:
            template_parent = re.sub(r"[①②③④⑤⑥⑦⑧⑨⑩\d]+$", "", skill["primary"]).strip()
            canonical = SPECIALTY_ALIASES.get(norm(template_parent), template_parent)
            if norm(canonical) == norm(wanted_parent) and skill["occupation_cell"] not in claimed:
                candidates.append(skill)
        preferred = [s for s in candidates if norm(s.get("specialty")) == norm(specialty)]
        if preferred:
            return preferred[0], None
        if candidates:
            return candidates[0], specialty

    partial = [s for s in skills if q and q in norm(s["name"]) and s["occupation_cell"] not in claimed]
    if len(partial) == 1:
        return partial[0], None
    raise PlanError(f"人物卡技能区找不到唯一技能槽位“{query}”；请先运行 catalog 查看可用名称")


def related_skill_slots(occupation: dict[str, Any], skills: list[dict[str, Any]]) -> list[dict[str, Any]]:
    source = str(occupation.get("skills") or "")
    chunks = re.split(r"[，,。；;、]", source)
    related: list[dict[str, Any]] = []
    seen: set[str] = set()
    for chunk in chunks:
        if not chunk.strip():
            continue
        try:
            skill, _ = resolve_skill(chunk.strip(), skills, set())
        except PlanError:
            continue
        if skill["name"] not in seen:
            related.append({"name": skill["name"], "occupation_cell": skill["occupation_cell"],
                            "interest_cell": skill["interest_cell"], "specialty_cell": skill["specialty_cell"]})
            seen.add(skill["name"])
    for skill in skills:
        if skill["name"] in seen:
            continue
        primary = re.sub(r"[①②③④⑤⑥⑦⑧⑨⑩\d]+$", "", skill["primary"]).rstrip("：:")
        source_norm = norm(source)
        direct = len(norm(primary)) >= 2 and norm(primary) in source_norm
        alias = any(norm(alias) in source_norm and norm(canonical) == norm(primary)
                    for alias, canonical in SPECIALTY_ALIASES.items())
        prefix = primary == "图书馆使用" and norm("图书馆") in source_norm
        if direct or alias or prefix:
            related.append({"name": skill["name"], "occupation_cell": skill["occupation_cell"],
                            "interest_cell": skill["interest_cell"], "specialty_cell": skill["specialty_cell"]})
            seen.add(skill["name"])
    return related


def compact_catalog(catalog: dict[str, Any], occupation_query: str | None,
                    only: str | None, weapon_query: str | None) -> dict[str, Any]:
    if only == "all":
        return catalog
    occupation = None
    candidates = []
    if occupation_query:
        query = norm(occupation_query)
        exact = [item for item in catalog["occupations"] if norm(item["name"]) == query]
        candidates = exact or [item for item in catalog["occupations"] if query in norm(item["name"]) or norm(item["name"]) in query]
        if len(candidates) == 1:
            occupation = candidates[0]
    result: dict[str, Any] = {"template": catalog["template"]}
    if only in (None, "occupation"):
        if occupation:
            result["occupation_match"] = occupation
            result["open_occupation_slots"] = open_occupation_slots(occupation.get("skills"))
        elif occupation_query and candidates:
            result["occupation_candidates"] = [
                {"name": item["name"], "serial": item["serial"], "skills": item["skills"],
                 "description": str(item["description"] or "")[:350],
                 "open_occupation_slots": open_occupation_slots(item.get("skills"))}
                for item in candidates[:8]
            ]
        elif occupation_query:
            result["occupation_candidates"] = []
        else:
            result["occupation_summary"] = {"count": len(catalog["occupations"]),
                                            "hint": "用 --occupation 名称查询；用 --all 查看完整目录"}
    if only in (None, "skills"):
        if occupation:
            result["skill_slots"] = related_skill_slots(occupation, catalog["skills"])
            result["open_occupation_slots"] = open_occupation_slots(occupation.get("skills"))
        elif only == "skills":
            result["skill_slots"] = [
                {"name": s["name"], "occupation_cell": s["occupation_cell"], "interest_cell": s["interest_cell"]}
                for s in catalog["skills"]]
        else:
            result["skill_summary"] = {"count": len(catalog["skills"]),
                                       "hint": "用 --occupation 名称查询相关槽位；用 --only skills 查看全部技能槽位"}
    if only in (None, "weapons"):
        if weapon_query:
            result["weapon_candidates"] = weapon_candidates(weapon_query, catalog["weapons"])
        else:
            examples = []
            for token in ("小型刀具", "小型棍状物", "小型自动手枪", "步枪", "霰弹枪"):
                match = next((w["name"] for w in catalog["weapons"] if token in w["name"]), None)
                if match:
                    examples.append(match)
            result["weapon_summary"] = {"count": len(catalog["weapons"]), "examples": examples,
                                        "hint": "用 --only weapons --weapon 名称查询候选"}
    return result


def formula_value(formula: Any, attrs: dict[str, int]) -> int:
    if isinstance(formula, (int, float)):
        return int(formula)
    expr = str(formula or "").strip().lstrip("=").upper().replace("；", ",").replace("，", ",")
    if not expr:
        raise PlanError("职业点公式为空")
    expr = re.sub(r"\bMAX\s*\(", "max(", expr)
    expr = re.sub(r"\bMIN\s*\(", "min(", expr)
    for key, value in attrs.items():
        expr = re.sub(rf"\b{re.escape(key)}\b", str(value), expr)
    tree = ast.parse(expr, mode="eval")
    allowed_bin = (ast.Add, ast.Sub, ast.Mult, ast.Div, ast.FloorDiv)
    for node in ast.walk(tree):
        if isinstance(node, ast.Expression | ast.Load | ast.Constant):
            continue
        if isinstance(node, ast.BinOp) and isinstance(node.op, allowed_bin):
            continue
        if isinstance(node, allowed_bin):
            continue
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
            continue
        if isinstance(node, (ast.UAdd, ast.USub)):
            continue
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id in {"max", "min"}:
            continue
        if isinstance(node, ast.Name) and node.id in {"max", "min"}:
            continue
        raise PlanError(f"不支持的职业点公式：{formula}")
    return int(eval(compile(tree, "<occupation-formula>", "eval"), {"__builtins__": {}}, {"max": max, "min": min}))


def _distribute_reduction(attrs: dict[str, int], keys: list[str], amount: int) -> dict[str, int]:
    remaining = amount
    reductions = {key: 0 for key in keys}
    while remaining > 0:
        key = max(keys, key=lambda item: (attrs[item] - reductions[item], -keys.index(item)))
        if attrs[key] - reductions[key] <= 0:
            break
        step = min(5, remaining, attrs[key] - reductions[key])
        reductions[key] += step
        remaining -= step
    if remaining:
        raise PlanError("年龄减值超过相关属性可承受范围")
    for key, value in reductions.items():
        attrs[key] -= value
    return {key: value for key, value in reductions.items() if value}


def age_adjust(raw: dict[str, int], age: int, rng: random.Random,
               override: dict[str, Any] | None, characteristics_are_final: bool = False) -> tuple[dict[str, int], dict[str, Any]]:
    attrs = dict(raw)
    log: dict[str, Any] = {"age": age, "education_rolls": [], "attribute_reductions": {}, "app_reduction": 0, "mov_penalty": 0}
    if not 15 <= age <= 89:
        raise PlanError("基础规则仅支持 15–89 岁调查员")

    if age <= 19:
        edu_rolls, reduction, app_reduction, mov_penalty, group = 0, 5, 0, 0, ["STR", "SIZ"]
    elif age <= 39:
        edu_rolls, reduction, app_reduction, mov_penalty, group = 1, 0, 0, 0, []
    elif age <= 49:
        edu_rolls, reduction, app_reduction, mov_penalty, group = 2, 5, 5, 1, ["STR", "CON", "DEX"]
    elif age <= 59:
        edu_rolls, reduction, app_reduction, mov_penalty, group = 3, 10, 10, 2, ["STR", "CON", "DEX"]
    elif age <= 69:
        edu_rolls, reduction, app_reduction, mov_penalty, group = 4, 20, 15, 3, ["STR", "CON", "DEX"]
    elif age <= 79:
        edu_rolls, reduction, app_reduction, mov_penalty, group = 4, 40, 20, 4, ["STR", "CON", "DEX"]
    else:
        edu_rolls, reduction, app_reduction, mov_penalty, group = 4, 80, 25, 5, ["STR", "CON", "DEX"]

    if characteristics_are_final:
        if override:
            raise PlanError("最终属性模式不能同时提供 age_adjustment")
        log["mode"] = "supplied_final"
        log["mov_penalty"] = mov_penalty
        return attrs, log

    log["mode"] = "adjusted_from_raw"
    override = override or {}
    requested_reductions = override.get("attribute_reductions")
    if requested_reductions is not None:
        parsed = {str(k).upper(): as_int(v, f"age_adjustment.attribute_reductions.{k}") for k, v in requested_reductions.items()}
        if set(parsed) - set(group) or sum(parsed.values()) != reduction or any(v < 0 for v in parsed.values()):
            raise PlanError(f"年龄属性减值必须只分配给 {group}，且合计为 {reduction}")
        for key, value in parsed.items():
            attrs[key] -= value
        log["attribute_reductions"] = parsed
    elif reduction:
        log["attribute_reductions"] = _distribute_reduction(attrs, group, reduction)

    attrs["APP"] = max(0, attrs["APP"] - app_reduction)
    log["app_reduction"] = app_reduction
    log["mov_penalty"] = mov_penalty

    if age <= 19:
        attrs["EDU"] = max(0, attrs["EDU"] - 5)
        second = override.get("luck_second_roll")
        second = as_int(second, "age_adjustment.luck_second_roll") if second is not None else rng.randint(1, 100)
        attrs["LUCK"] = max(attrs["LUCK"], second)
        log["luck_second_roll"] = second
    else:
        supplied_rolls = override.get("education_rolls")
        if supplied_rolls is not None and len(supplied_rolls) != edu_rolls:
            raise PlanError(f"该年龄需要 {edu_rolls} 次 EDU 增强检定")
        for index in range(edu_rolls):
            roll = as_int(supplied_rolls[index], f"education_rolls[{index}]") if supplied_rolls is not None else rng.randint(1, 100)
            success = roll > attrs["EDU"]
            gain = rng.randint(1, 10) if success else 0
            before = attrs["EDU"]
            attrs["EDU"] = min(99, attrs["EDU"] + gain)
            log["education_rolls"].append({"roll": roll, "before": before, "success": success, "gain": attrs["EDU"] - before})

    if any(value < 0 or value > 100 for value in attrs.values()):
        raise PlanError("年龄调整后的属性必须在 0–100 范围内")
    return attrs, log


def derived(attrs: dict[str, int], age_log: dict[str, Any]) -> dict[str, Any]:
    total = attrs["STR"] + attrs["SIZ"]
    bands = [
        (64, "-2", -2), (84, "-1", -1), (124, "0", 0),
        (164, "+1D4", 1), (204, "+1D6", 2), (284, "+2D6", 3),
        (364, "+3D6", 4), (444, "+4D6", 5), (524, "+5D6", 6),
    ]
    db, build = "+6D6", 7
    for ceiling, candidate_db, candidate_build in bands:
        if total <= ceiling:
            db, build = candidate_db, candidate_build
            break
    if attrs["STR"] < attrs["SIZ"] and attrs["DEX"] < attrs["SIZ"]:
        mov = 7
    elif attrs["STR"] > attrs["SIZ"] and attrs["DEX"] > attrs["SIZ"]:
        mov = 9
    else:
        mov = 8
    mov = max(1, mov - int(age_log["mov_penalty"]))
    return {
        "HP": math.floor((attrs["CON"] + attrs["SIZ"]) / 10),
        "MP": math.floor(attrs["POW"] / 5), "SAN": attrs["POW"],
        "DB": db, "Build": build, "MOV": mov,
        "Dodge_base": math.floor(attrs["DEX"] / 2), "Language_base": attrs["EDU"],
    }


def sheet_xml_path(xlsx: Path, sheet_name: str) -> str:
    with zipfile.ZipFile(xlsx) as zf:
        workbook = ET.fromstring(zf.read("xl/workbook.xml"))
        rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    rid = None
    for sheet in workbook.findall(f"{{{NS_MAIN}}}sheets/{{{NS_MAIN}}}sheet"):
        if sheet.get("name") == sheet_name:
            rid = sheet.get(f"{{{NS_REL_DOC}}}id")
            break
    if not rid:
        raise PlanError(f"工作簿找不到页签 {sheet_name}")
    target = None
    for rel in rels.findall(f"{{{NS_REL_PKG}}}Relationship"):
        if rel.get("Id") == rid:
            target = rel.get("Target")
            break
    if not target:
        raise PlanError(f"无法解析页签 {sheet_name} 的 OOXML 路径")
    target = target.replace("\\", "/")
    if target.startswith("/"):
        return target.lstrip("/")
    return "xl/" + target.lstrip("./")


def column_number(cell_ref: str) -> int:
    letters = re.match(r"[A-Z]+", cell_ref.upper()).group(0)
    value = 0
    for char in letters:
        value = value * 26 + ord(char) - 64
    return value


def set_cell(root: ET.Element, ref: str, value: Any) -> None:
    sheet_data = root.find(f"{{{NS_MAIN}}}sheetData")
    row_no = int(re.search(r"\d+", ref).group(0))
    row = sheet_data.find(f"{{{NS_MAIN}}}row[@r='{row_no}']")
    if row is None:
        row = ET.SubElement(sheet_data, f"{{{NS_MAIN}}}row", {"r": str(row_no)})
    cell = row.find(f"{{{NS_MAIN}}}c[@r='{ref}']")
    if cell is None:
        cell = ET.Element(f"{{{NS_MAIN}}}c", {"r": ref})
        before = next((c for c in row.findall(f"{{{NS_MAIN}}}c") if column_number(c.get("r")) > column_number(ref)), None)
        if before is None:
            row.append(cell)
        else:
            row.insert(list(row).index(before), cell)
    for child in list(cell):
        if child.tag in {f"{{{NS_MAIN}}}f", f"{{{NS_MAIN}}}v", f"{{{NS_MAIN}}}is"}:
            cell.remove(child)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        cell.attrib.pop("t", None)
        ET.SubElement(cell, f"{{{NS_MAIN}}}v").text = str(value)
    else:
        cell.set("t", "inlineStr")
        inline = ET.SubElement(cell, f"{{{NS_MAIN}}}is")
        text = ET.SubElement(inline, f"{{{NS_MAIN}}}t")
        rendered = str(value)
        if rendered.startswith(" ") or rendered.endswith(" ") or "\n" in rendered:
            text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        text.text = rendered


def patch_xlsx(template: Path, output: Path, values: dict[str, Any]) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    sheet_path = sheet_xml_path(template, MAIN_SHEET)
    with zipfile.ZipFile(template, "r") as zin:
        sheet_root = ET.fromstring(zin.read(sheet_path))
        workbook_root = ET.fromstring(zin.read("xl/workbook.xml"))
        for ref, value in values.items():
            set_cell(sheet_root, ref, value)
        calc_pr = workbook_root.find(f"{{{NS_MAIN}}}calcPr")
        if calc_pr is None:
            calc_pr = ET.SubElement(workbook_root, f"{{{NS_MAIN}}}calcPr")
        calc_pr.set("calcMode", "auto")
        calc_pr.set("fullCalcOnLoad", "1")
        calc_pr.set("forceFullCalc", "1")
        modified = {
            sheet_path: ET.tostring(sheet_root, encoding="utf-8", xml_declaration=True),
            "xl/workbook.xml": ET.tostring(workbook_root, encoding="utf-8", xml_declaration=True),
        }
        fd, temp_name = tempfile.mkstemp(suffix=".xlsx", dir=output.parent)
        os.close(fd)
        Path(temp_name).unlink(missing_ok=True)
        try:
            with zipfile.ZipFile(temp_name, "w") as zout:
                for item in zin.infolist():
                    zout.writestr(item, modified.get(item.filename, zin.read(item.filename)))
            shutil.move(temp_name, output)
        finally:
            Path(temp_name).unlink(missing_ok=True)


def inspect_cells(workbook: Path, sheet_name: str, refs: list[str]) -> dict[str, Any]:
    normalized_refs = [str(ref).strip().upper() for ref in refs]
    if not normalized_refs or any(not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", ref) for ref in normalized_refs):
        raise PlanError("--cells 请传入逗号分隔的单元格地址，例如 E3,E5,AA61")
    sheet_path = sheet_xml_path(workbook, sheet_name)
    with zipfile.ZipFile(workbook) as zf:
        root = ET.fromstring(zf.read(sheet_path))
        shared = []
        if "xl/sharedStrings.xml" in zf.namelist():
            strings_root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
            shared = ["".join(si.itertext()) for si in strings_root.findall(f"{{{NS_MAIN}}}si")]
    requested = set(normalized_refs)
    found: dict[str, Any] = {ref: None for ref in normalized_refs}
    for cell in root.iter(f"{{{NS_MAIN}}}c"):
        ref = cell.get("r", "").upper()
        if ref not in requested:
            continue
        formula = cell.find(f"{{{NS_MAIN}}}f")
        value_node = cell.find(f"{{{NS_MAIN}}}v")
        inline = cell.find(f"{{{NS_MAIN}}}is")
        raw_value = value_node.text if value_node is not None else None
        if inline is not None:
            value: Any = "".join(inline.itertext())
        elif cell.get("t") == "s" and raw_value is not None:
            value = shared[int(raw_value)]
        elif cell.get("t") == "b" and raw_value is not None:
            value = raw_value == "1"
        elif raw_value is not None and cell.get("t") not in {"str", "e"}:
            number = float(raw_value)
            value = int(number) if number.is_integer() else number
        else:
            value = raw_value
        found[ref] = {"value": value, "formula": formula.text if formula is not None else None}
    return {"workbook": str(workbook), "sheet": sheet_name, "cells": found}


def normalize_characteristics(data: dict[str, Any]) -> dict[str, int]:
    upper = {str(k).upper(): v for k, v in data.items()}
    missing = [key for key in CHAR_KEYS if key not in upper]
    if missing:
        raise PlanError("缺少属性：" + "、".join(missing))
    attrs = {key: as_int(upper[key], f"characteristics.{key}") for key in CHAR_KEYS}
    if any(value < 1 or value > 100 for value in attrs.values()):
        raise PlanError("原始属性必须在 1–100 范围内")
    return attrs


def fill(template: Path, plan_path: Path, output: Path) -> dict[str, Any]:
    plan = json.loads(plan_path.read_text(encoding="utf-8-sig"))
    basic = plan.get("basic") or {}
    background = plan.get("background") or {}
    required_basic = ("name", "age", "era", "occupation", "gender")
    missing_basic = [key for key in required_basic if basic.get(key) in (None, "")]
    if missing_basic:
        raise PlanError("basic 缺少：" + "、".join(missing_basic))
    if output.suffix.lower() != ".xlsx":
        output = output / output_filename(basic["name"], basic.get("group_name", ""))
    missing_background = [key for key in BACKGROUND_CELLS if not str(background.get(key, "")).strip()]
    if missing_background:
        raise PlanError("background 缺少：" + "、".join(missing_background))

    raw = normalize_characteristics(plan.get("characteristics") or {})
    characteristics_are_final = plan.get("characteristics_are_final", False)
    if not isinstance(characteristics_are_final, bool):
        raise PlanError("characteristics_are_final 必须是 true 或 false")
    age = as_int(basic["age"], "basic.age")
    seed_text = str(plan.get("random_seed") or hashlib.sha256(json.dumps(plan, ensure_ascii=False, sort_keys=True).encode()).hexdigest())
    rng = random.Random(seed_text)
    attrs, age_log = age_adjust(raw, age, rng, plan.get("age_adjustment"), characteristics_are_final)
    catalog = workbook_catalog(template)
    occupation = resolve_named(str(basic["occupation"]), catalog["occupations"], "职业")
    occ_pool = formula_value(occupation["points_formula"], attrs)
    interest_pool = attrs["INT"] * 2

    values: dict[str, Any] = {
        "E3": basic["name"], "M3": basic.get("player", ""), "E4": age,
        "M4": basic["era"], "E5": occupation["name"], "M5": occupation["serial"],
        "E6": basic["gender"], "M6": basic.get("nationality", ""),
        "E7": basic.get("residence", ""), "M7": basic.get("hometown", ""),
    }
    date = basic.get("date") or {}
    if date.get("year") not in (None, ""):
        values["G8"] = as_int(date["year"], "basic.date.year")
    if date.get("month") not in (None, ""):
        values["J8"] = f"{as_int(date['month'], 'basic.date.month')}月"
    if date.get("day") not in (None, ""):
        values["L8"] = f"{as_int(date['day'], 'basic.date.day')}日"
    if date.get("time") not in (None, ""):
        values["N8"] = date["time"]
    for key, cell in CHAR_CELLS.items():
        values[cell] = attrs[key]
    for key, cell in BACKGROUND_CELLS.items():
        values[cell] = str(background[key]).strip()

    claimed: set[str] = set()
    skill_report = []
    selected_open_skills = []
    occ_total = int_total = 0
    for index, entry in enumerate(plan.get("skills") or []):
        if not entry.get("name"):
            raise PlanError(f"skills[{index}] 缺少 name")
        occ = as_int(entry.get("occupation", 0), f"skills[{index}].occupation")
        interest = as_int(entry.get("interest", 0), f"skills[{index}].interest")
        if occ < 0 or interest < 0 or occ + interest == 0:
            raise PlanError(f"skills[{index}] 的投入点必须为非负且不能全为 0")
        skill, specialty = resolve_skill(str(entry["name"]), catalog["skills"], claimed)
        claimed.add(skill["occupation_cell"])
        if specialty:
            values[skill["specialty_cell"]] = specialty
        if occ:
            values[skill["occupation_cell"]] = occ
        if interest:
            values[skill["interest_cell"]] = interest
        occ_total += occ
        int_total += interest
        if entry.get("occupation_choice") and not occ:
            raise PlanError(f"skills[{index}] 标记为职业开放名额，但没有分配职业点")
        if entry.get("occupation_choice"):
            if skill["primary"] not in catalog["open_choice_names"]:
                raise PlanError(f"人物卡任意特长下拉表中没有“{skill['primary']}”")
            selected_open_skills.append(skill["primary"])
        skill_report.append({"requested": entry["name"], "slot": skill["name"], "specialty_written": specialty,
                             "occupation": occ, "interest": interest, "occupation_choice": bool(entry.get("occupation_choice"))})

    open_slots = open_occupation_slots(occupation.get("skills"))
    open_count = sum(item["count"] for item in open_slots)
    chosen_count = sum(bool(item["occupation_choice"]) for item in skill_report)
    if open_count and chosen_count != open_count:
        raise PlanError(f"职业说明含 {open_count} 个开放技能名额，计划标记了 {chosen_count} 个；请在相应 skills 项设置 occupation_choice: true")
    if not open_count and chosen_count:
        raise PlanError("该职业说明没有可识别的开放技能名额，不应标记 occupation_choice")
    if len(selected_open_skills) > len(OPEN_CHOICE_CELLS):
        raise PlanError(f"任意特长超过模板 {len(OPEN_CHOICE_CELLS)} 个可计数槽位")
    for cell, skill_name in zip(OPEN_CHOICE_CELLS, selected_open_skills):
        values[cell] = skill_name

    if occ_total != occ_pool:
        raise PlanError(f"职业技能点合计 {occ_total}，应为 {occ_pool}")
    if int_total != interest_pool:
        raise PlanError(f"兴趣技能点合计 {int_total}，应为 {interest_pool}")
    credit = next((row["occupation"] for row in skill_report if norm(row["requested"]) == norm("信用评级")), 0)
    credit_min = as_int(occupation.get("credit_min") or 0, "职业信用评级下限")
    credit_max = as_int(occupation.get("credit_max") or 99, "职业信用评级上限")
    if not credit_min <= credit <= credit_max:
        raise PlanError(f"信用评级职业点为 {credit}，职业范围是 {credit_min}–{credit_max}")

    carry_row = backpack_row = 79
    item_report = []
    for index, item in enumerate(plan.get("items") or []):
        name = str(item.get("name", "")).strip()
        if not name:
            raise PlanError(f"items[{index}] 缺少 name")
        location = str(item.get("location") or "").strip()
        packed = bool(item.get("backpack")) or any(bag in location for bag in ("背包", "挎包", "提包", "行李箱"))
        for part in (piece.strip() for piece in re.split(r"[、，,]", name)):
            if not part:
                continue
            if packed:
                if backpack_row > 93:
                    raise PlanError("背包物品超过模板 15 行")
                values[f"N{backpack_row}"] = part
                item_report.append({"name": part, "cell": f"N{backpack_row}"})
                backpack_row += 1
            else:
                if carry_row > 93:
                    raise PlanError("随身物品超过模板 15 行")
                values[f"F{carry_row}"] = part
                if location:
                    values[f"D{carry_row}"] = location
                if item.get("status"):
                    values[f"B{carry_row}"] = item["status"]
                item_report.append({"name": part, "cell": f"F{carry_row}"})
                carry_row += 1

    weapon_report = []
    for offset, requested in enumerate(plan.get("weapons") or []):
        if offset >= 5:
            raise PlanError("武器超过模板新增槽位 5 件")
        weapon = resolve_named(str(requested), catalog["weapons"], "武器")
        row = 54 + offset
        values[f"B{row}"] = weapon["name"]
        values[f"G{row}"] = weapon["name"]
        weapon_report.append({"requested": requested, "matched": weapon["name"], "row": row})

    if output.exists():
        raise PlanError(f"输出文件已存在，不覆盖：{output}")
    patch_xlsx(template, output, values)
    report = {
        "output": str(output), "seed": seed_text,
        "occupation": occupation, "raw_characteristics": raw,
        "final_characteristics": attrs, "characteristics_are_final": characteristics_are_final,
        "age_adjustment": age_log,
        "derived": derived(attrs, age_log),
        "points": {"occupation_pool": occ_pool, "occupation_used": occ_total, "interest_pool": interest_pool, "interest_used": int_total, "credit_rating": credit},
        "open_occupation_slots": open_slots, "selected_open_skills": selected_open_skills,
        "skills": skill_report, "items": item_report, "weapons": weapon_report,
        "written_cells": values,
    }
    report_path = output.with_suffix(".report.json")
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description="Inspect or fill the bundled CoC 7e investigator workbook")
    sub = parser.add_subparsers(dest="command", required=True)
    catalog_parser = sub.add_parser("catalog", help="Show a compact occupation and related template slots")
    catalog_parser.add_argument("--template", type=Path, required=True)
    catalog_parser.add_argument("--occupation")
    catalog_parser.add_argument("--only", choices=("occupation", "skills", "weapons"))
    catalog_parser.add_argument("--weapon", help="Show up to five likely weapon names")
    catalog_parser.add_argument("--all", action="store_true", help="Print the complete template catalog")
    fill_parser = sub.add_parser("fill", help="Validate a plan and create a filled workbook")
    fill_parser.add_argument("--template", type=Path, required=True)
    fill_parser.add_argument("--input", type=Path, required=True)
    fill_parser.add_argument("--output", type=Path, required=True)
    inspect_parser = sub.add_parser("inspect", help="Read selected cells from the OOXML workbook as UTF-8 JSON")
    inspect_parser.add_argument("--workbook", type=Path, required=True)
    inspect_parser.add_argument("--sheet", default=MAIN_SHEET)
    inspect_parser.add_argument("--cells", required=True)
    args = parser.parse_args()
    try:
        if args.command == "catalog":
            data = workbook_catalog(args.template)
            result = data if args.all else compact_catalog(data, args.occupation, args.only, args.weapon)
            print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
        elif args.command == "fill":
            report = fill(args.template, args.input, args.output)
            print(json.dumps({"output": report["output"], "points": report["points"], "derived": report["derived"]}, ensure_ascii=False, indent=2))
        else:
            result = inspect_cells(args.workbook, args.sheet, args.cells.split(","))
            print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except (PlanError, json.JSONDecodeError, OSError, zipfile.BadZipFile) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
