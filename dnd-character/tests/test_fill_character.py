import copy
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from fill_character import PlanError, fill, validate_plan  # noqa: E402


BASE = {
    "kind": "dnd-character-plan", "schema_version": 1,
    "basic": {"name": "测试角色", "class": "战士", "species": "人类", "background": "士兵",
              "level": 1, "group_name": "单元测试", "story": "来自边境村落，受训后加入巡逻队。"},
    "scores_base": {"STR": 15, "DEX": 13, "CON": 14, "INT": 10, "WIS": 12, "CHA": 8},
    "save_proficiencies": ["STR", "CON"], "skill_proficiencies": ["运动", "察觉"],
    "inventory": ["背包", "火把"], "weapons": ["长剑"], "languages": ["通用语"],
}


def five_r_plan():
    plan = copy.deepcopy(BASE)
    plan.update(ruleset="5r", background_bonus={"STR": 2, "CON": 1},
                inventory=[{"name": "背包", "quantity": 1, "unit_weight": 5, "location": "bag1"},
                           {"name": "火把", "quantity": 2, "unit_weight": 1, "location": "bag1"}],
                feats=[{"name": "起源专长测试", "type": "origin", "acquired_level": 1,
                        "source": "背景", "description": "完整的专长说明"}],
                combat={"initiative_bonus": 0, "ac_bonus": 0,
                        "shield": {"name": "", "equipped": False}, "spell_slots_current": {},
                        "resources": [], "currency": {"gp": 0, "pp": 0, "ep": 0, "sp": 0, "cp": 0}})
    return plan


class CharacterFillTests(unittest.TestCase):
    def test_5r_preserves_template_structure_and_writes_origin(self):
        plan = five_r_plan()
        plan.update(class_features=[{"name": "战斗风格", "description": "测试描述"}],
                    species_traits=[{"name": "人类特性", "description": "测试特性"}],
                    known_spells=["测试法术"],
                    prepared_spells=[{"name": "测试法术", "source": "职业"}], cantrips=["测试戏法"])
        plan["combat"]["spell_slots_current"] = {"1": 1}
        with tempfile.TemporaryDirectory() as temporary:
            destination = fill(plan, Path(temporary), None)
            source = ROOT / "assets" / "dnd5r-blank.xlsx"
            with zipfile.ZipFile(source) as original, zipfile.ZipFile(destination) as result:
                self.assertIsNone(result.testzip())
                self.assertEqual(original.namelist(), result.namelist())
                self.assertEqual(original.read("xl/worksheets/sheet3.xml"), result.read("xl/worksheets/sheet3.xml"))
            wb = openpyxl.load_workbook(destination, read_only=True)
            self.assertEqual(wb["主要"]["E3"].value, "测试角色")
            self.assertEqual(wb["主要"]["K13"].value, 2)
            self.assertEqual(wb["主要"]["K15"].value, 1)
            self.assertEqual(wb["主要"]["B41"].value, "O")
            self.assertEqual(wb["起源"]["S17"].value, plan["basic"]["story"])
            self.assertEqual(wb["主要"]["AX17"].value, "战斗风格")
            self.assertEqual(wb["主要"]["BT6"].value, "人类特性")
            self.assertIsNone(wb["主要"]["Q66"].value)
            self.assertEqual(wb["主要"]["Y66"].value, "测试法术")
            self.assertEqual(wb["主要"]["Q77"].value, "测试戏法")
            self.assertEqual(wb["主要"]["BT19"].value, 1)
            self.assertEqual(wb["背包"]["Y5"].value, "背包")
            self.assertEqual(wb["背包"]["AP5"].value, 5)
            self.assertEqual(wb["背包"]["AS6"].value, 2)
            self.assertIsNone(wb["背包"]["B5"].value)
            self.assertEqual(wb["法术书"]["X3"].value, "测试法术")
            self.assertEqual(wb["主要"]["BF7"].data_type, "f")
            source_wb = openpyxl.load_workbook(source, read_only=True)
            self.assertEqual(wb.sheetnames, source_wb.sheetnames)
            source_wb.close()
            wb.close()

    def test_5e_has_distinct_racial_bonus_fields(self):
        plan = copy.deepcopy(BASE)
        plan.update(ruleset="5e", racial_bonus={"STR": 1})
        plan["basic"].update(hit_die=10, armor_type="重甲", armor="锁子甲", speed="30尺",
                             background_feature="军旅身份")
        with tempfile.TemporaryDirectory() as temporary:
            destination = fill(plan, Path(temporary), None)
            with zipfile.ZipFile(destination) as archive:
                self.assertIsNone(archive.testzip())
            wb = openpyxl.load_workbook(destination, read_only=True)
            sh = wb["主要"]
            self.assertEqual(sh["K14"].value, 1)
            self.assertEqual(sh["AW4"].value, 10)
            self.assertEqual(sh["A14"].value, True)
            self.assertEqual(sh["F14"].data_type, "f")
            self.assertEqual(wb["背景"]["U12"].value, "军旅身份")
            self.assertEqual(wb["背景"]["U18"].value, plan["basic"]["story"])
            self.assertEqual(len(wb.sheetnames), 19)
            wb.close()

    def test_rejects_mixed_edition_ability_bonus(self):
        plan = five_r_plan()
        plan["racial_bonus"] = {"STR": 2}
        with self.assertRaisesRegex(PlanError, "不使用 racial_bonus"):
            validate_plan(plan)

    def test_5r_spells_inventory_feats_and_combat_inputs(self):
        plan = five_r_plan()
        plan["basic"].update(level=8, armor="链甲", **{"class": "圣武士"})
        plan["feats"].extend([
            {"name": "警觉", "type": "general", "acquired_level": 4, "source": "等级提升", "description": "完整警觉说明" * 10},
            {"name": "防御", "type": "fighting_style", "acquired_level": 8, "source": "职业", "description": "完整防御说明"},
        ])
        plan["inventory"].extend([
            {"name": "绳索", "quantity": 1, "unit_weight": 10, "location": "bag2"},
            {"name": "长剑", "quantity": 1, "unit_weight": 3, "location": "equipped"},
            {"name": "箭矢", "quantity": 20, "unit_weight": 0.05, "location": "ammo"},
        ])
        plan["free_cast_spells"] = [{"name": "至圣斩", "source": "职业特性", "uses_max": 1,
                                     "uses_remaining": 1, "recharge": "长休", "description": "完整法术说明"}]
        plan["prepared_spells"] = [{"name": f"常备法术{i}", "source": "圣武士"} for i in range(1, 8)]
        plan["combat"].update(initiative_bonus=3, initiative_bonus_source="警觉",
                              ac_bonus=1, ac_bonus_source="防御",
                              shield={"name": "盾牌", "equipped": True},
                              spell_slots_current={"1": 2},
                              resources=[{"name": "圣疗", "current": 5, "max": 5, "recharge": "长休"}],
                              currency={"gp": 10, "pp": 1, "ep": 0, "sp": 2, "cp": 3})
        with tempfile.TemporaryDirectory() as temporary:
            destination = fill(plan, Path(temporary), None)
            wb = openpyxl.load_workbook(destination, read_only=True)
            main, bag = wb["主要"], wb["背包"]
            self.assertEqual(main["Q66"].value, "至圣斩")
            self.assertEqual([main[cell].value for cell in ("Y66", "AG66", "AO66", "Y67", "AG67", "AO67", "Y68")],
                             [f"常备法术{i}" for i in range(1, 8)])
            self.assertEqual([main[f"BT{r}"].value for r in (19, 20, 21)], [1, 4, 8])
            self.assertLessEqual(len(main["BZ20"].value), 32)
            self.assertEqual((main["G22"].value, main["G23"].value), (3, 1))
            self.assertIn("G22", main["D22"].value)
            self.assertIn("AS40", main["D23"].value)
            self.assertEqual(main["L40"].value, "链甲")
            self.assertEqual((main["AP40"].value, main["AS40"].value), ("盾牌", "是"))
            self.assertEqual((main["J26"].value, main["M26"].value), (5, 5))
            self.assertEqual((main["J27"].value, main["M27"].value), (1, 1))
            self.assertEqual((main["AK30"].value, main["AS30"].value), ("箭矢", 20))
            self.assertEqual(main["E72"].value, 2)
            self.assertEqual((main["AI59"].value, main["AP59"].value, main["AI61"].value), (10, 1, 2))
            self.assertEqual((bag["B29"].value, bag["S29"].value, bag["V29"].value), ("绳索", 10, 1))
            self.assertEqual(main["P60"].value, "=背包!AS3")
            self.assertEqual(main["T60"].value, "=背包!V27")
            self.assertIsNone(bag["B5"].value)
            wb.close()
            reference = destination.with_name(destination.stem + "-战斗速查.md")
            self.assertIn("完整警觉说明" * 10, reference.read_text(encoding="utf-8"))
            self.assertIn("完整法术说明", reference.read_text(encoding="utf-8"))
            report = json.loads(destination.with_suffix(".report.json").read_text(encoding="utf-8"))
            self.assertEqual(report["bag_weight_lb"], {"bag1": 7, "bag2": 10})
            self.assertEqual(report["native_recalculation"], "pending")

    def test_5r_rejects_missing_weight_and_feat_level(self):
        plan = five_r_plan()
        del plan["inventory"][0]["unit_weight"]
        with self.assertRaisesRegex(PlanError, "unit_weight"):
            validate_plan(plan)
        plan = five_r_plan()
        del plan["feats"][0]["acquired_level"]
        with self.assertRaisesRegex(PlanError, "acquired_level"):
            validate_plan(plan)
        plan = five_r_plan()
        plan["inventory"].append({"name": "长剑", "quantity": 1, "unit_weight": 3, "location": "bag1"})
        with self.assertRaisesRegex(PlanError, "additional_copy"):
            validate_plan(plan)
        plan = five_r_plan()
        plan["prepared_spells"] = [{"name": "测试法术", "source": "职业"}]
        with self.assertRaisesRegex(PlanError, "spell_slots_current"):
            validate_plan(plan)

    def test_5e_features_spells_and_level_hp(self):
        plan = copy.deepcopy(BASE)
        plan.update(ruleset="5e", racial_bonus={"STR": 1},
                    class_features=[{"name": "战斗风格", "description": "测试描述"}],
                    species_traits=[{"name": "额外熟练", "description": "测试种族特性"}],
                    known_spells=["测试法术"])
        plan["basic"].update(level=4, hit_die=10, subclass="勇士")
        with tempfile.TemporaryDirectory() as temporary:
            destination = fill(plan, Path(temporary), None)
            wb = openpyxl.load_workbook(destination, read_only=True)
            sh = wb["主要"]
            self.assertEqual(sh["M29"].value, 36)
            self.assertEqual(sh["L32"].value, 4)
            self.assertEqual(sh["C90"].value, "子职：勇士")
            self.assertEqual(sh["C91"].value, "战斗风格")
            self.assertEqual(sh["Z90"].value, "额外熟练")
            self.assertEqual(sh["C72"].value, "测试法术")
            self.assertEqual(sh["B70"].value, "已知法术")
            self.assertEqual(wb["背包"]["B5"].value, "背包")
            wb.close()


if __name__ == "__main__":
    unittest.main()
